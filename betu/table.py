from .helpers import tr

"""Tabular import and plotting, independent of the user interface."""
import csv
import io
from datetime import date, datetime
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
import numpy as np
from ._plotting import (
    DEFAULT_COLORS,
    DEFAULT_LINEWIDTHS,
    DEFAULT_MARKERS,
    DEFAULT_MARKERSIZES,
    apply_grid,
    apply_legend,
    line_styles,
    fill_text_color,
)
from .helpers import CHART_LABELS

CHART_TYPES = {pair[0]: key for key, pair in CHART_LABELS.items()}


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return (
            value.isoformat(sep=" ")
            if isinstance(value, datetime)
            else value.isoformat()
        )
    return str(value)


def has_header(rows):
    """Only infer a header when text is followed by numeric observations."""

    def is_value(value):
        value = cell_text(value).strip()
        try:
            float(value)
            return True
        except ValueError:
            try:
                datetime.fromisoformat(value)
                return True
            except ValueError:
                return False

    if len(rows) < 2 or any(is_value(value) for value in rows[0]):
        return False
    for col, name in enumerate(rows[0]):
        if not cell_text(name).strip():
            continue
        values = [
            row[col]
            for row in rows[1:21]
            if col < len(row) and cell_text(row[col]).strip()
        ]
        if values and sum(is_value(value) for value in values) > len(values) / 2:
            return True
    return False


def normalize_table(rows, header=True):
    rows = [list(row) for row in rows]
    while rows and all((cell_text(v).strip() == "" for v in rows[-1])):
        rows.pop()
    if not rows:
        raise ValueError(tr("message_15"))
    width = max(map(len, rows))
    while width and all(
        (len(row) < width or not cell_text(row[width - 1]).strip() for row in rows)
    ):
        width -= 1
    if not width:
        raise ValueError(tr("message_16"))
    rows = [
        [cell_text(v) for v in row[:width]] + [""] * max(0, width - len(row))
        for row in rows
    ]
    if header is None:
        header = has_header(rows)
    names = (
        rows.pop(0) if header else [tr("message_17", v0=i + 1) for i in range(width)]
    )
    headers = []
    for i, name in enumerate(names):
        base = name.strip() or tr("message_18", v0=i + 1)
        name, suffix = (base, 2)
        while name in headers:
            name = f"{base}_{suffix}"
            suffix += 1
        headers.append(name)
    return (headers, rows)


def parse_pasted_table(text, header=True):
    if not text.strip():
        raise ValueError(tr("message_19"))
    sample = text[:8192]
    if "\t" in sample:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;").delimiter
        except csv.Error:
            delimiter = ","
    return normalize_table(csv.reader(io.StringIO(text), delimiter=delimiter), header)


def excel_sheets(path):
    if Path(path).suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_table(path, sheet=None, header=True):
    """Read .xlsx/.xlsm/.xls or UTF-8/GB18030 CSV/TSV; never modify the input."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        return parse_pasted_table(text, header)
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            ws = (
                book.sheet_by_name(sheet)
                if isinstance(sheet, str)
                else book.sheet_by_index(sheet or 0)
            )
            rows = []
            for i in range(ws.nrows):
                row = []
                for cell in ws.row(i):
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    row.append(value)
                rows.append(row)
            return normalize_table(rows, header)
        finally:
            book.release_resources()
    if suffix not in (".xlsx", ".xlsm"):
        raise ValueError(tr("message_20"))
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet] if isinstance(sheet, str) else book.worksheets[sheet or 0]
        return normalize_table(ws.iter_rows(values_only=True), header)
    finally:
        book.close()


def _numeric(values, name):
    result = []
    for row, value in enumerate(values, 2):
        if value is None or cell_text(value).strip() == "":
            result.append(np.nan)
            continue
        try:
            number = float(value)
            if not np.isfinite(number):
                raise ValueError
        except (TypeError, ValueError):
            raise ValueError(tr("message_21", v0=row, v1=name, v2=value)) from None
        result.append(number)
    return np.asarray(result)


def plot_table(
    data,
    x=None,
    series=None,
    kind="line",
    *,
    bins=10,
    x_name=None,
    y_name=None,
    save_dir=None,
    location="best",
    ncol=1,
    fsize=14,
    figsize=[6, 5],
    xt_rotation=0,
    xrotation=0,
    yrotation=0,
    linestyles=None,
    linewidth=DEFAULT_LINEWIDTHS,
    markers=DEFAULT_MARKERS,
    markersize=DEFAULT_MARKERSIZES,
    colors=DEFAULT_COLORS,
    isgrid=True,
    xpad=3,
    ypad=3,
    xlabelpad=10,
    ylabelpad=10,
    xlabelsize="auto",
    ylabelsize="auto",
    legendsize="auto",
    legend_options=None,
    yt_rotation=0,
    show_legend=True,
):
    """Plot named columns. Blank numeric cells are missing, never replaced by zero.

    line/bar/barh: one observation per row; box: optional x grouping;
    pie: exactly one nonnegative series; hist: independent numeric samples;
    scatter: numeric x required. Returns matplotlib.pyplot.
    Pie percentage labels automatically use black or white for fill contrast.
    """
    kind = CHART_TYPES.get(kind, kind)
    if kind not in CHART_TYPES.values():
        raise ValueError(tr("message_22"))
    if (
        not data
        or len({len(v) for v in data.values()}) != 1
        or (not len(next(iter(data.values()))))
    ):
        raise ValueError(tr("message_23"))
    if x is not None and x not in data:
        raise ValueError(tr("message_24"))
    series = list(series) if series is not None else [k for k in data if k != x]
    if (
        not series
        or len(set(series)) != len(series)
        or any((k not in data or k == x for k in series))
    ):
        raise ValueError(tr("message_25"))
    if kind == "pie" and len(series) != 1:
        raise ValueError(tr("message_26"))
    if kind == "scatter" and x is None:
        raise ValueError(tr("message_27"))
    if kind == "hist" and (not isinstance(bins, int) or bins < 1):
        raise ValueError(tr("message_28"))
    values = {name: _numeric(data[name], name) for name in series}
    for name, ys in values.items():
        if not np.isfinite(ys).any():
            raise ValueError(tr("message_29", v0=name))
    count = len(next(iter(data.values())))
    labels = (
        [cell_text(v) for v in data[x]]
        if x is not None
        else [str(i + 1) for i in range(count)]
    )
    xs = np.arange(count, dtype=float)
    numeric_x = False
    if kind in ("line", "scatter") and x is not None:
        try:
            xs = _numeric(data[x], x)
            numeric_x = True
        except ValueError:
            if kind == "scatter":
                raise
    if kind == "pie":
        ys = values[series[0]]
        if (ys[np.isfinite(ys)] < 0).any() or np.nansum(ys) <= 0:
            raise ValueError(tr("message_30"))
    if kind in ("line", "scatter"):
        for name, ys in values.items():
            if not (np.isfinite(xs) & np.isfinite(ys)).any():
                raise ValueError(tr("message_31", v0=name))
    fig, ax = plt.subplots(figsize=figsize)
    styles = list(
        line_styles(len(series), linestyles, linewidth, markers, markersize, colors)
    )
    try:
        handles = None
        if kind == "pie":
            ys = values[series[0]]
            valid = np.isfinite(ys)
            palette = [
                style[-1]
                for style in line_styles(
                    int(valid.sum()), linestyles, linewidth, markers, markersize, colors
                )
            ]
            handles, _, percentages = ax.pie(
                ys[valid],
                colors=palette,
                labels=None,
                autopct="%1.1f%%",
                textprops={"fontsize": fsize},
            )
            for wedge, label in zip(handles, np.asarray(labels)[valid]):
                wedge.set_label(label)
            for wedge, text in zip(handles, percentages):
                text.set_color(
                    fill_text_color(wedge.get_facecolor(), fig.get_facecolor())
                )
            ax.axis("equal")
        elif kind == "box":
            handles = []
            groups = list(dict.fromkeys(labels)) if x is not None else [None]
            width = 0.7 / len(series)
            for j, name in enumerate(series):
                ys = values[name]
                color = styles[j][-1]
                samples, positions = ([], [])
                for i, group in enumerate(groups):
                    valid = np.isfinite(ys)
                    if group is not None:
                        valid &= np.asarray(labels) == group
                    if valid.any():
                        samples.append(ys[valid])
                        positions.append(
                            i + (j - (len(series) - 1) / 2) * width
                            if x is not None
                            else j
                        )
                if samples:
                    artists = ax.boxplot(
                        samples,
                        positions=positions,
                        widths=width if x is not None else 0.6,
                        patch_artist=True,
                        manage_ticks=False,
                    )
                    for box in artists["boxes"]:
                        box.set_facecolor(color)
                        box.set_alpha(0.65)
                handles.append(Patch(facecolor=color, alpha=0.65, label=name))
            ax.set_xticks(range(len(groups) if x is not None else len(series)))
            ax.set_xticklabels(groups if x is not None else series)
        else:
            for j, name in enumerate(series):
                ys = values[name]
                ls, lw, marker, ms, color = styles[j]
                if kind == "line":
                    ax.plot(
                        xs,
                        ys,
                        label=name,
                        color=color,
                        linestyle=ls,
                        linewidth=lw,
                        marker=marker,
                        markersize=ms,
                        markevery=max(1, count // 10),
                    )
                elif kind == "scatter":
                    valid = np.isfinite(xs) & np.isfinite(ys)
                    ax.scatter(
                        xs[valid],
                        ys[valid],
                        label=name,
                        color=color,
                        marker=marker or "o",
                        s=ms**2,
                    )
                elif kind in ("bar", "barh"):
                    width = 0.8 / len(series)
                    pos = xs + (j - (len(series) - 1) / 2) * width
                    if kind == "bar":
                        ax.bar(pos, ys, width=width, label=name, color=color)
                    else:
                        ax.barh(pos, ys, height=width, label=name, color=color)
                elif kind == "hist":
                    ax.hist(
                        ys[np.isfinite(ys)],
                        bins=bins,
                        label=name,
                        color=color,
                        alpha=0.65,
                        edgecolor="white",
                    )
            if kind == "barh":
                ax.set_yticks(xs)
                ax.set_yticklabels(labels)
            elif kind in ("bar", "line") and (not numeric_x):
                step = max(1, count // 12)
                ax.set_xticks(xs[::step])
                ax.set_xticklabels(labels[::step])
        if kind != "pie":
            default_x = (
                tr("message_32") if kind in ("hist", "barh") else x or tr("message_33")
            )
            default_y = (
                x or tr("message_34")
                if kind == "barh"
                else tr("message_35") if kind == "hist" else tr("message_36")
            )
            ax.set_xlabel(
                default_x if x_name is None else x_name,
                fontsize=fsize if xlabelsize == "auto" else xlabelsize,
                rotation=xrotation,
                labelpad=xlabelpad,
            )
            ax.set_ylabel(
                default_y if y_name is None else y_name,
                fontsize=fsize if ylabelsize == "auto" else ylabelsize,
                rotation=yrotation,
                labelpad=ylabelpad,
                horizontalalignment="right" if yrotation % 360 == 0 else "center",
            )
            ax.tick_params(
                axis="x", pad=xpad, labelsize=fsize, labelrotation=xt_rotation
            )
            ax.tick_params(
                axis="y", pad=ypad, labelsize=fsize, labelrotation=yt_rotation
            )
            apply_grid(ax, isgrid)
        apply_legend(
            ax,
            location,
            ncol,
            fsize,
            legendsize,
            legend_options if show_legend else False,
            handles=handles,
        )
        fig.tight_layout()
        if save_dir is not None:
            fig.savefig(save_dir, dpi=300, bbox_inches="tight")
        return plt
    except Exception:
        plt.close(fig)
        raise
